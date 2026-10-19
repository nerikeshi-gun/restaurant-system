import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models import DiningTable, MenuItem, Message, Order, ServiceRequest, Store, TableSession
from app.seed import seed_database


@pytest.fixture()
def shared_client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_database(db)

    def override_db():
        with Session(engine) as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            yield client, engine
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(engine)


def test_staff_table_uses_dedicated_handheld_page(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session = db.scalar(select(TableSession).order_by(TableSession.id))
        table_id = session.table_id
        table_number = db.get(DiningTable, table_id).number

    assert client.get("/staff").status_code == 200
    response = client.get(f"/staff/table/{table_id}")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert 'class="staff-handheld"' in response.text
    assert "音声注文" not in response.text
    assert "/static/staff-order.js" in response.text
    assert 'id="confirm-dialog"' in response.text
    assert 'id="submit-confirm"' in response.text

    language_page = client.get(f"/table/{table_number}")
    assert language_page.headers["cache-control"] == "no-store"
    assert "言語を選んでください" in language_page.text
    assert f'/table/{table_number}/order?lang=en' in language_page.text
    assert f'/table/{table_number}/order?lang=ur' in language_page.text
    assert "اردو" in language_page.text

    tablet_page = client.get(f"/table/{table_number}/order?lang=en")
    assert tablet_page.headers["cache-control"] == "no-store"
    assert 'id="customer-confirm-dialog"' in tablet_page.text
    assert 'id="submit-customer-confirm"' in tablet_page.text
    assert 'id="start-over"' in tablet_page.text
    assert "最初から" in tablet_page.text
    assert "注文へ</button>" not in tablet_page.text
    assert 'id="start-voice"' in tablet_page.text
    assert 'id="start-menu"' in tablet_page.text
    assert "注文方法を選択してください" in tablet_page.text
    assert "食べたいもの・飲みたいものと数量を話してください" in tablet_page.text
    assert "customer-order-page" in tablet_page.text
    assert 'data-language="en"' in tablet_page.text
    assert 'data-order-mode="choice"' in tablet_page.text
    assert "order.js?v=20260828-7" in tablet_page.text
    assert "注文候補を確認" in tablet_page.text
    assert "話した内容を確認</button>" not in tablet_page.text
    assert 'data-request="water"' not in tablet_page.text
    assert "ご注文内容" in tablet_page.text
    assert "このテーブルの注文" not in tablet_page.text

    voice_page = client.get(f"/table/{table_number}/order?lang=en&mode=voice")
    assert 'order-mode-voice' in voice_page.text
    assert f'data-menu-url="/table/{table_number}/order?lang=en&amp;mode=menu"' in voice_page.text
    menu_page = client.get(f"/table/{table_number}/order?lang=en&mode=menu")
    assert 'order-mode-menu' in menu_page.text
    assert f'data-start-url="/table/{table_number}/order?lang=en"' in menu_page.text
    urdu_page = client.get(f"/table/{table_number}/order?lang=ur&mode=voice")
    assert '<html lang="ur" dir="rtl">' in urdu_page.text
    assert 'data-language="ur"' in urdu_page.text


def test_staff_and_customer_receive_the_same_draft_and_kitchen_status(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
        menu_item_id = db.scalar(select(MenuItem.id).where(MenuItem.name == "Aランチ"))

    initial = client.get(f"/api/v1/sessions/{session_id}/shared-order").json()
    order_id = initial["id"]
    with client.websocket_connect(f"/api/v1/ws/orders/{session_id}") as staff_ws:
        with client.websocket_connect(f"/api/v1/ws/orders/{session_id}") as customer_ws:
            response = client.post(
                f"/api/v1/shared-orders/{order_id}/items",
                json={
                    "menu_item_id": menu_item_id,
                    "quantity": 2,
                    "modifier_option_ids": [],
                    "input_by": "staff",
                    "input_language": "ja",
                    "original_text": "Aランチを二つ",
                    "expected_version": initial["version"],
                },
            )
            assert response.status_code == 200
            staff_event = staff_ws.receive_json()
            customer_event = customer_ws.receive_json()
            assert staff_event["data"] == customer_event["data"]
            assert customer_event["data"]["items"][0]["input_by"] == "staff"
            assert customer_event["data"]["items"][0]["status"] == "draft"

            draft = response.json()
            confirmed = client.post(
                f"/api/v1/shared-orders/{order_id}/confirm",
                json={
                    "expected_version": draft["version"],
                    "confirmed_by": "customer",
                    "confirmation_method": "touch_button",
                },
            )
            assert confirmed.status_code == 200
            assert confirmed.json()["display_number"] == 1
            assert confirmed.json()["items"][0]["status"] == "confirmed"
            assert confirmed.json()["items"][0]["confirmed_at"] is not None
            assert confirmed.json()["pending_total"] == 0
            assert confirmed.json()["current_total"] == confirmed.json()["items"][0]["line_total"]
            assert confirmed.json()["pos_export_preview"]["adapter"] == "json_mock"
            staff_ws.receive_json()
            customer_ws.receive_json()

            kitchen = client.get("/api/v1/kitchen/snapshot").json()
            line = kitchen["shared_orders"][0]["items"][0]
            assert line["quantity"] == 2
            assert line["kitchen_status"] == "not_started"

            changed = client.patch(
                f"/api/v1/kitchen/items/{line['id']}/status",
                json={"status": "preparing"},
            )
            assert changed.status_code == 200
            assert staff_ws.receive_json()["data"]["items"][0]["kitchen_status"] == "preparing"
            assert customer_ws.receive_json()["data"]["items"][0]["kitchen_status"] == "preparing"

            customer_cancel = client.post(
                f"/api/v1/shared-orders/{order_id}/items/{line['id']}/cancel",
                json={"expected_version": changed.json()["version"], "input_by": "customer"},
            )
            assert customer_cancel.status_code == 200
            assert customer_cancel.json()["items"][0]["status"] == "cancel_requested"
            assert staff_ws.receive_json()["data"]["items"][0]["status"] == "cancel_requested"
            assert customer_ws.receive_json()["data"]["items"][0]["status"] == "cancel_requested"
            kitchen_cancel = client.patch(
                f"/api/v1/kitchen/items/{line['id']}/status",
                json={"status": "cancelled"},
            )
            assert kitchen_cancel.status_code == 409
            rejected = client.post(f"/api/v1/kitchen/items/{line['id']}/cancel-request/reject")
            assert rejected.status_code == 200
            assert rejected.json()["items"][0]["status"] == "confirmed"
            staff_ws.receive_json()
            customer_ws.receive_json()

            served = client.patch(
                f"/api/v1/kitchen/items/{line['id']}/status",
                json={"status": "served"},
            )
            assert served.status_code == 200
            staff_ws.receive_json()
            customer_ws.receive_json()
            served_snapshot = client.get("/api/v1/kitchen/snapshot").json()["shared_orders"]
            assert served_snapshot[0]["items"][0]["kitchen_status"] == "served"


def test_stale_version_cannot_overwrite_shared_order(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
        menu_item_id = db.scalar(select(MenuItem.id).where(MenuItem.name == "アジフライ定食"))
    order = client.get(f"/api/v1/sessions/{session_id}/shared-order").json()
    payload = {
        "menu_item_id": menu_item_id,
        "quantity": 1,
        "modifier_option_ids": [],
        "input_by": "customer",
        "input_language": "ja",
        "expected_version": order["version"],
    }
    assert client.post(f"/api/v1/shared-orders/{order['id']}/items", json=payload).status_code == 200
    conflict = client.post(f"/api/v1/shared-orders/{order['id']}/items", json=payload)
    assert conflict.status_code == 409
    assert "更新" in conflict.json()["detail"]


def test_old_business_day_order_is_not_reused(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
    old_order = client.get(f"/api/v1/sessions/{session_id}/shared-order").json()
    with Session(engine) as db:
        order = db.get(Order, old_order["id"])
        order.business_date = "2000-01-01"
        db.commit()
    current_order = client.get(f"/api/v1/sessions/{session_id}/shared-order").json()
    assert current_order["id"] != old_order["id"]


def test_store_name_color_and_order_screen_layout_are_configurable(shared_client):
    client, engine = shared_client
    initial = client.get("/api/v1/admin/store-settings")
    assert initial.status_code == 200
    assert initial.json()["layout"]["desktop"]["block_order"] == ["input", "menu", "service"]
    assert initial.json()["layout"]["tablet"]["block_order"] == ["input", "menu", "service"]

    updated = client.patch(
        "/api/v1/admin/store-settings",
        json={
            "name": "テスト食堂",
            "theme_color": "#315F86",
            "layout": {
                "desktop": {"block_order": ["service", "input", "menu"], "order_position": "right"},
                "tablet": {"block_order": ["menu", "service", "input"], "order_position": "top"},
            },
        },
    )
    assert updated.status_code == 200
    assert updated.json()["theme_color"] == "#315f86"

    settings = client.get("/api/v1/admin/store-settings").json()
    assert settings["name"] == "テスト食堂"
    assert settings["layout"]["desktop"]["block_order"] == ["service", "input", "menu"]
    assert settings["layout"]["tablet"]["block_order"] == ["menu", "service", "input"]
    assert settings["layout"]["tablet"]["order_position"] == "top"

    language_page = client.get("/table/1")
    assert language_page.status_code == 200
    assert "テスト食堂" in language_page.text
    assert "--accent:#315f86" in language_page.text

    page = client.get("/table/1/order?lang=ja")
    assert page.status_code == 200
    assert "テスト食堂" in page.text
    assert "--accent:#315f86" in page.text
    assert '"order_position": "top"' in page.text

    with Session(engine) as db:
        assert db.get(Store, "STORE001").name == "テスト食堂"


def test_store_layout_rejects_duplicate_boxes(shared_client):
    client, _ = shared_client
    response = client.patch(
        "/api/v1/admin/store-settings",
        json={
            "name": "テスト食堂",
            "theme_color": "#315f86",
            "layout": {
                "desktop": {"block_order": ["input", "input", "service"], "order_position": "right"},
                "tablet": {"block_order": ["input", "menu", "service"], "order_position": "top"},
            },
        },
    )
    assert response.status_code == 422


def test_store_layout_change_is_pushed_to_open_order_screens(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
    payload = {
        "name": "同期テスト食堂",
        "theme_color": "#27815c",
        "layout": {
            "desktop": {"block_order": ["service", "menu", "input"], "order_position": "bottom"},
            "tablet": {"block_order": ["menu", "input", "service"], "order_position": "top"},
        },
    }
    with client.websocket_connect(f"/api/v1/ws/orders/{session_id}") as websocket:
        assert client.patch("/api/v1/admin/store-settings", json=payload).status_code == 200
        event = websocket.receive_json()
    assert event["type"] == "store.settings_changed"
    assert event["data"]["layout"] == payload["layout"]
    assert event["data"]["theme_color"] == "#27815c"


def test_kitchen_reply_is_pushed_and_can_be_reloaded(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))

    with client.websocket_connect(f"/api/v1/ws/orders/{session_id}") as customer_ws:
        response = client.post(
            "/api/v1/messages",
            json={"session_id": session_id, "sender_type": "kitchen", "text": "あと5分ほどです"},
        )
        assert response.status_code == 200
        event = customer_ws.receive_json()
        assert event["type"] == "message.created"
        assert event["data"]["text"] == "あと5分ほどです"

    latest = client.get(f"/api/v1/sessions/{session_id}/messages/latest")
    assert latest.status_code == 200
    assert latest.headers["cache-control"] == "no-store"
    assert latest.json()["id"] == response.json()["id"]
    assert latest.json()["text"] == "あと5分ほどです"


def test_dismissed_kitchen_reply_is_not_reloaded_and_dismissal_is_pushed(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_ids = list(db.scalars(select(TableSession.id).order_by(TableSession.id)).all())
    session_id = session_ids[0]

    older = client.post(
        "/api/v1/messages",
        json={"session_id": session_id, "sender_type": "kitchen", "text": "先ほどの返答"},
    ).json()
    latest = client.post(
        "/api/v1/messages",
        json={"session_id": session_id, "sender_type": "kitchen", "text": "最新の返答"},
    ).json()

    with client.websocket_connect(f"/api/v1/ws/orders/{session_id}") as customer_ws:
        dismissed = client.patch(
            f"/api/v1/sessions/{session_id}/messages/{latest['id']}/dismiss"
        )
        assert dismissed.status_code == 200
        event = customer_ws.receive_json()
        assert event["type"] == "message.dismissed"
        assert event["data"]["id"] == latest["id"]

    assert client.get(f"/api/v1/sessions/{session_id}/messages/latest").json() is None
    assert client.patch(
        f"/api/v1/sessions/{session_id}/messages/{latest['id']}/dismiss"
    ).status_code == 200

    other_session_id = session_ids[1]
    assert client.patch(
        f"/api/v1/sessions/{other_session_id}/messages/{older['id']}/dismiss"
    ).status_code == 404


def test_admin_can_manage_categories_and_assign_menu_items(shared_client):
    client, _ = shared_client
    created = client.post("/api/v1/admin/menu-categories", json={"name": "おすすめ", "display_order": 0})
    assert created.status_code == 200
    category_id = created.json()["id"]

    menu = client.get("/api/v1/admin/menu").json()
    item_id = menu[0]["id"]
    assigned = client.patch(f"/api/v1/admin/menu/{item_id}/category", json={"category_id": category_id})
    assert assigned.status_code == 200
    assert assigned.json()["category"] == "おすすめ"

    renamed = client.patch(f"/api/v1/admin/menu-categories/{category_id}", json={"name": "店長おすすめ", "display_order": 1})
    assert renamed.status_code == 200
    assert renamed.json()["name"] == "店長おすすめ"
    assert client.delete(f"/api/v1/admin/menu-categories/{category_id}").status_code == 409

    empty = client.post("/api/v1/admin/menu-categories", json={"name": "削除用", "display_order": 99}).json()
    assert client.delete(f"/api/v1/admin/menu-categories/{empty['id']}").status_code == 200


def test_admin_can_register_and_edit_menu_item(shared_client):
    client, _ = shared_client
    category_id = client.get("/api/v1/admin/menu-categories").json()[0]["id"]
    created = client.post(
        "/api/v1/admin/menu",
        json={
            "menu_code": "M900",
            "category_id": category_id,
            "name": "テスト定食",
            "short_name": "テスト",
            "reading": "てすとていしょく",
            "aliases": ["試験定食"],
            "localized_names": {"en": "Test Set Meal", "zh": "测试套餐"},
            "price": 980,
            "display_order": 3,
            "size_label": "並",
            "set_contents": "ご飯・味噌汁",
            "is_available": True,
        },
    )
    assert created.status_code == 200
    item_id = created.json()["id"]
    visible_item = next(item for item in client.get("/api/v1/menu").json() if item["id"] == item_id)
    assert visible_item["localized_names"]["en"] == "Test Set Meal"
    assert visible_item["localized_names"]["zh"] == "测试套餐"

    payload = created.json() | {"name": "テスト定食 改", "price": 1080, "aliases": ["試験定食", "確認定食"]}
    payload.pop("id")
    updated = client.put(f"/api/v1/admin/menu/{item_id}", json=payload)
    assert updated.status_code == 200
    assert updated.json()["name"] == "テスト定食 改"
    assert updated.json()["price"] == 1080
    assert updated.json()["aliases"] == ["試験定食", "確認定食"]

    duplicate = client.post("/api/v1/admin/menu", json=payload)
    assert duplicate.status_code == 409

    deleted = client.delete(f"/api/v1/admin/menu/{item_id}")
    assert deleted.status_code == 200
    assert all(item["id"] != item_id for item in client.get("/api/v1/admin/menu").json())
    assert all(item["id"] != item_id for item in client.get("/api/v1/menu").json())
    assert client.delete(f"/api/v1/admin/menu/{item_id}").status_code == 404


def test_menu_template_download_and_simple_csv_import(shared_client):
    client, _ = shared_client
    template = client.get("/api/v1/admin/menu-template.csv")
    assert template.status_code == 200
    assert "attachment" in template.headers["content-disposition"]
    assert template.content.startswith(b"\xef\xbb\xbf")
    template_text = template.content.decode("utf-8-sig")
    assert "menu_id,name_ja,price,category,category_order,available,alias,reading,short_name,display_order,size,set_contents,toppings,localized_names" in template_text
    assert '"辛さ,大,中,小"' in template_text

    current = client.get("/api/v1/admin/menu-current.csv")
    assert current.status_code == 200
    assert "attachment" in current.headers["content-disposition"]
    assert current.content.startswith(b"\xef\xbb\xbf")
    current_text = current.content.decode("utf-8-sig")
    assert "Aランチ" in current_text
    assert '"ビールサイズ,小,中,大"' in current_text

    simple_csv = (
        "menu_id,name_ja,price,category,available,alias\n"
        "M901,チキンカレー,1100,カレー,1,チキン|チキンカレー\n"
        "M902,マトンマサラ,1000,カレー,0,マトン\n"
    ).encode("utf-8-sig")
    imported = client.post(
        "/api/v1/admin/menu/import-csv",
        files={"file": ("menu.csv", simple_csv, "text/csv")},
    )
    assert imported.status_code == 200
    assert imported.json()["store_id"] == "STORE001"
    assert imported.json()["items"] == 2
    admin_menu = client.get("/api/v1/admin/menu").json()
    assert next(item for item in admin_menu if item["menu_code"] == "M901")["is_available"] is True
    assert next(item for item in admin_menu if item["menu_code"] == "M902")["is_available"] is False


def test_required_modifier_is_managed_and_validated_without_llm(shared_client):
    client, engine = shared_client
    group = client.post(
        "/api/v1/admin/modifier-groups",
        json={"name": "辛さテスト", "selection_type": "single", "display_order": 1, "is_active": True},
    )
    assert group.status_code == 200
    group_id = group.json()["id"]
    updated_group = client.patch(
        f"/api/v1/admin/modifier-groups/{group_id}",
        json={"name": "辛さテスト更新", "selection_type": "single", "display_order": 2, "is_active": True},
    )
    assert updated_group.status_code == 200
    assert updated_group.json()["name"] == "辛さテスト更新"
    assert updated_group.json()["display_order"] == 2
    option = client.post(
        f"/api/v1/admin/modifier-groups/{group_id}/options",
        json={"name": "中辛", "price_delta": 50, "is_default": True, "display_order": 1, "is_available": True},
    )
    assert option.status_code == 200
    option_id = option.json()["id"]

    admin_menu = client.get("/api/v1/admin/menu").json()
    menu_id = next(item["id"] for item in admin_menu if item["menu_code"] == "M001")
    payload = client.get(f"/api/v1/admin/menu/{menu_id}").json()
    payload.pop("id")
    payload["modifier_groups"] = [{"group_id": group_id, "required": True, "display_order": 1}]
    assert client.put(f"/api/v1/admin/menu/{menu_id}", json=payload).status_code == 200

    visible = next(item for item in client.get("/api/v1/menu").json() if item["id"] == menu_id)
    assert visible["modifiers"][0]["name"] == "辛さテスト更新"
    assert visible["modifiers"][0]["required"] is True

    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
    order = client.get(f"/api/v1/sessions/{session_id}/shared-order").json()
    base = {
        "menu_item_id": menu_id,
        "quantity": 1,
        "input_by": "customer",
        "input_language": "ja",
        "expected_version": order["version"],
    }
    missing = client.post(f"/api/v1/shared-orders/{order['id']}/items", json=base | {"modifier_option_ids": []})
    assert missing.status_code == 422
    assert "辛さテスト更新" in missing.json()["detail"]

    accepted = client.post(f"/api/v1/shared-orders/{order['id']}/items", json=base | {"modifier_option_ids": [option_id]})
    assert accepted.status_code == 200
    line = accepted.json()["items"][0]
    assert line["modifiers"] == ["辛さテスト更新：中辛"]
    assert line["unit_price"] == payload["price"] + 50
    assert client.delete(f"/api/v1/admin/modifier-options/{option_id}").status_code == 409
    assert client.delete(f"/api/v1/admin/modifier-groups/{group_id}").status_code == 409


def test_unused_option_button_and_group_can_be_deleted(shared_client):
    client, _ = shared_client
    group = client.post(
        "/api/v1/admin/modifier-groups",
        json={"name": "削除テスト", "selection_type": "single", "display_order": 999, "is_active": True},
    ).json()
    option = client.post(
        f"/api/v1/admin/modifier-groups/{group['id']}/options",
        json={"name": "削除するボタン", "price_delta": 0, "is_default": False, "display_order": 1, "is_available": True},
    ).json()
    assert client.delete(f"/api/v1/admin/modifier-options/{option['id']}").status_code == 200
    assert client.delete(f"/api/v1/admin/modifier-groups/{group['id']}").status_code == 200
    assert client.delete(f"/api/v1/admin/modifier-groups/{group['id']}").status_code == 404


def test_simple_toppings_csv_creates_groups_and_options(shared_client):
    client, _ = shared_client
    csv_data = (
        "menu_id,name_ja,price,category,available,toppings\n"
        'M950,テストカレー,900,カレー,1,"辛さ,大,中,小 / ご飯量,大盛,普通,小盛"\n'
    ).encode("utf-8-sig")
    imported = client.post(
        "/api/v1/admin/menu/import-csv",
        files={"file": ("menu.csv", csv_data, "text/csv")},
    )
    assert imported.status_code == 200
    groups = client.get("/api/v1/admin/modifier-groups").json()
    assert [option["name"] for option in next(group for group in groups if group["name"] == "辛さ")["options"]] == ["大", "中", "小"]
    menu_item = next(item for item in client.get("/api/v1/menu").json() if item["name"] == "テストカレー")
    assert [group["name"] for group in menu_item["modifiers"]] == ["辛さ", "ご飯量"]
    assert next(group for group in menu_item["modifiers"] if group["name"] == "辛さ")["required"] is True
    assert next(group for group in menu_item["modifiers"] if group["name"] == "ご飯量")["required"] is False


def test_excel_menu_sheet_has_group_dropdown_and_can_be_imported(shared_client):
    client, _ = shared_client
    template = client.get("/api/v1/admin/menu-template.xlsx")
    assert template.status_code == 200
    workbook = load_workbook(io.BytesIO(template.content))
    sheet = workbook["メニュー"]
    assert sheet["M1"].value == "選択肢グループ1"
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    assert validations[0].formula1 == "=ToppingGroups"

    current = client.get("/api/v1/admin/menu-current.xlsx")
    current_book = load_workbook(io.BytesIO(current.content))
    current_sheet = current_book["メニュー"]
    rows = list(current_sheet.iter_rows(min_row=2, values_only=True))
    beer = next(row for row in rows if row[1] == "キリン一番搾り 生ビール")
    assert "ビールサイズ" in beer[12:17]

    sheet["A2"] = "M951"
    sheet["B2"] = "Excelテスト商品"
    sheet["C2"] = 1200
    sheet["D2"] = "料理"
    sheet["M2"] = "盛り"
    output = io.BytesIO()
    workbook.save(output)
    imported = client.post(
        "/api/v1/admin/menu/import-csv",
        files={"file": ("menu.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200
    menu_item = next(item for item in client.get("/api/v1/menu").json() if item["name"] == "Excelテスト商品")
    assert [group["name"] for group in menu_item["modifiers"]] == ["盛り"]


@pytest.mark.parametrize("language", ["ja", "en", "zh", "ko", "vi", "ur"])
def test_customer_language_and_order_mode_route_matrix(shared_client, language):
    client, _ = shared_client
    choice = client.get(f"/table/1/order?lang={language}")
    assert choice.status_code == 200
    assert f'data-language="{language}"' in choice.text
    assert 'data-order-mode="choice"' in choice.text
    for mode in ["voice", "menu"]:
        page = client.get(f"/table/1/order?lang={language}&mode={mode}")
        assert page.status_code == 200
        assert f'order-mode-{mode}' in page.text
    if language == "ur":
        assert 'dir="rtl"' in choice.text


def test_customer_routes_reject_unknown_resources_and_normalize_parameters(shared_client):
    client, _ = shared_client
    assert client.get("/table/999999").status_code == 404
    assert client.get("/table/999999/order?lang=ja").status_code == 404
    assert client.get("/t/not-a-token").status_code == 404
    assert client.get("/t/not-a-token/order?lang=ja").status_code == 404
    normalized = client.get("/table/1/order?lang=xx&mode=invalid")
    assert normalized.status_code == 200
    assert 'data-language="ja"' in normalized.text
    assert 'data-order-mode="choice"' in normalized.text


def test_shared_order_rejects_quantity_boundaries_missing_products_and_mismatched_candidates(shared_client):
    client, engine = shared_client
    initial = client.get("/api/v1/tables/1/shared-order").json()
    menu_item = client.get("/api/v1/menu").json()[0]
    base = {"menu_item_id": menu_item["id"], "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": initial["version"]}
    for quantity in [0, 100]:
        assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=base | {"quantity": quantity}).status_code == 422
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=base | {"menu_item_id": 999999, "quantity": 1}).status_code == 422
    mismatch = base | {"quantity": 1, "selected_candidate": {"menu_item_id": 999999}}
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=mismatch).status_code == 422
    with Session(engine) as db:
        db.get(MenuItem, menu_item["id"]).is_available = False
        db.commit()
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=base | {"quantity": 1}).status_code == 422


def test_confirm_cancel_and_kitchen_transition_failure_matrix(shared_client):
    client, _ = shared_client
    initial = client.get("/api/v1/tables/1/shared-order").json()
    order_id = initial["id"]
    empty_confirm = {"expected_version": initial["version"], "confirmed_by": "customer", "confirmation_method": "touch_button"}
    assert client.post(f"/api/v1/shared-orders/{order_id}/confirm", json=empty_confirm).status_code == 422
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    added = client.post(f"/api/v1/shared-orders/{order_id}/items", json={"menu_item_id": menu_item_id, "quantity": 1, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": initial["version"]}).json()
    confirmed = client.post(f"/api/v1/shared-orders/{order_id}/confirm", json={"expected_version": added["version"], "confirmed_by": "customer", "confirmation_method": "touch_button"}).json()
    item_id = confirmed["items"][0]["id"]
    requested = client.post(f"/api/v1/shared-orders/{order_id}/items/{item_id}/cancel", json={"expected_version": confirmed["version"], "input_by": "customer"})
    assert requested.status_code == 200
    assert requested.json()["items"][0]["status"] == "cancel_requested"
    kitchen_item = client.get("/api/v1/kitchen/snapshot").json()["shared_orders"][0]["items"][0]
    assert kitchen_item["status"] == "cancel_requested"
    rejected = client.post(f"/api/v1/kitchen/items/{item_id}/cancel-request/reject")
    assert rejected.status_code == 200
    assert rejected.json()["items"][0]["status"] == "confirmed"
    assert client.patch(f"/api/v1/kitchen/items/{item_id}/status", json={"status": "preparing"}).status_code == 200
    current = client.get("/api/v1/tables/1/shared-order").json()
    late_cancel = client.post(f"/api/v1/shared-orders/{order_id}/items/{item_id}/cancel", json={"expected_version": current["version"], "input_by": "customer"})
    assert late_cancel.status_code == 200
    assert late_cancel.json()["items"][0]["status"] == "cancel_requested"
    assert client.post(f"/api/v1/kitchen/items/{item_id}/cancel-request/reject").status_code == 200
    assert client.patch(f"/api/v1/kitchen/items/{item_id}/status", json={"status": "not_started"}).status_code == 409
    assert client.patch(f"/api/v1/kitchen/items/{item_id}/status", json={"status": "served"}).status_code == 200
    assert client.patch(f"/api/v1/kitchen/items/{item_id}/status", json={"status": "preparing"}).status_code == 409


def test_approved_cancel_is_subtracted_from_customer_and_bill_totals(shared_client):
    client, _ = shared_client
    initial = client.get("/api/v1/tables/1/shared-order").json()
    order_id = initial["id"]
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    added = client.post(
        f"/api/v1/shared-orders/{order_id}/items",
        json={"menu_item_id": menu_item_id, "quantity": 2, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": initial["version"]},
    ).json()
    confirmed = client.post(
        f"/api/v1/shared-orders/{order_id}/confirm",
        json={"expected_version": added["version"], "confirmed_by": "customer", "confirmation_method": "touch_button"},
    ).json()
    item = confirmed["items"][0]
    original_total = item["line_total"]
    requested = client.post(
        f"/api/v1/shared-orders/{order_id}/items/{item['id']}/cancel",
        json={"expected_version": confirmed["version"], "input_by": "customer"},
    ).json()
    assert requested["current_total"] == original_total
    assert client.get(f"/api/v1/sessions/{requested['session_id']}/bill").json()["total"] == original_total
    approved = client.post(f"/api/v1/kitchen/items/{item['id']}/cancel-request/approve").json()
    assert approved["current_total"] == 0
    assert approved["confirmed_total"] == 0
    assert client.get(f"/api/v1/sessions/{approved['session_id']}/bill").json()["total"] == 0


def test_item_from_another_shared_order_cannot_be_changed_or_cancelled(shared_client):
    client, engine = shared_client
    first = client.get("/api/v1/tables/1/shared-order").json()
    second = client.get("/api/v1/tables/2/shared-order").json()
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    added = client.post(f"/api/v1/shared-orders/{first['id']}/items", json={"menu_item_id": menu_item_id, "quantity": 1, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": first["version"]}).json()
    foreign_item_id = added["items"][0]["id"]
    assert client.patch(f"/api/v1/shared-orders/{second['id']}/items/{foreign_item_id}", json={"quantity": 2, "expected_version": second["version"]}).status_code == 422
    assert client.post(f"/api/v1/shared-orders/{second['id']}/items/{foreign_item_id}/cancel", json={"expected_version": second["version"], "input_by": "staff"}).status_code == 422
    with Session(engine) as db:
        assert db.get(Order, second["id"]).version == second["version"]


def test_service_request_validation_actions_and_checkout_session_state(shared_client):
    client, engine = shared_client
    with Session(engine) as db:
        session_id = db.scalar(select(TableSession.id).order_by(TableSession.id))
    assert client.post("/api/v1/service-requests", json={"session_id": 999999, "request_type": "other", "message": "test"}).status_code == 404
    created = client.post("/api/v1/service-requests", json={"session_id": session_id, "request_type": "other", "message": "スプーン"})
    request_id = created.json()["id"]
    assert client.patch(f"/api/v1/service-requests/{request_id}/invalid").status_code == 400
    assert client.patch(f"/api/v1/service-requests/{request_id}/acknowledge").json()["status"] == "acknowledged"
    assert client.patch(f"/api/v1/service-requests/{request_id}/complete").json()["status"] == "completed"
    assert client.post("/api/v1/service-requests", json={"session_id": session_id, "request_type": "checkout", "message": ""}).status_code == 200
    with Session(engine) as db:
        assert db.get(TableSession, session_id).status == "checkout_requested"


def test_translation_and_shared_order_schema_validation_matrix(shared_client):
    client, _ = shared_client
    assert client.post("/api/v1/translation/prepare", json={"source_text": "", "source_language": "en"}).status_code == 422
    translation = client.post("/api/v1/translation/prepare", json={"source_text": "two curries", "source_language": "en"})
    assert translation.status_code == 200
    assert translation.json()["manual_confirmation_required"] is True
    assert translation.json()["translated_text"] == ""
    initial = client.get("/api/v1/tables/1/shared-order").json()
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    payload = {"menu_item_id": menu_item_id, "quantity": 1, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": initial["version"]}
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=payload | {"notes": "x" * 501}).status_code == 422
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=payload | {"input_by": "unknown"}).status_code == 422
    assert client.post(f"/api/v1/shared-orders/{initial['id']}/items", json=payload | {"expected_version": 0}).status_code == 422


def test_display_number_reset_boundaries_and_next_confirmation(shared_client):
    client, _ = shared_client
    assert client.post("/api/v1/admin/display-number/reset", json={"next_number": 0}).status_code == 422
    assert client.post("/api/v1/admin/display-number/reset", json={"next_number": 10000}).status_code == 422
    assert client.post("/api/v1/admin/display-number/reset", json={"next_number": 37}).json()["next_number"] == 37
    initial = client.get("/api/v1/tables/1/shared-order").json()
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    added = client.post(f"/api/v1/shared-orders/{initial['id']}/items", json={"menu_item_id": menu_item_id, "quantity": 1, "modifier_option_ids": [], "input_by": "staff", "input_language": "ja", "expected_version": initial["version"]}).json()
    confirmed = client.post(f"/api/v1/shared-orders/{initial['id']}/confirm", json={"expected_version": added["version"], "confirmed_by": "staff", "confirmation_method": "touch_button"})
    assert confirmed.status_code == 200
    assert confirmed.json()["display_number"] == 37


def test_admin_resets_departed_table_preserves_history_and_opens_fresh_session(shared_client):
    client, engine = shared_client
    initial = client.get("/api/v1/tables/1/shared-order").json()
    old_session_id = initial["session_id"]
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]

    first_added = client.post(
        f"/api/v1/shared-orders/{initial['id']}/items",
        json={"menu_item_id": menu_item_id, "quantity": 1, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": initial["version"]},
    ).json()
    confirmed = client.post(
        f"/api/v1/shared-orders/{initial['id']}/confirm",
        json={"expected_version": first_added["version"], "confirmed_by": "customer", "confirmation_method": "touch_button"},
    ).json()
    served = client.patch(
        f"/api/v1/kitchen/items/{confirmed['items'][0]['id']}/status",
        json={"status": "served"},
    ).json()
    with_draft = client.post(
        f"/api/v1/shared-orders/{initial['id']}/items",
        json={"menu_item_id": menu_item_id, "quantity": 2, "modifier_option_ids": [], "input_by": "customer", "input_language": "ja", "expected_version": served["version"]},
    ).json()
    service_request = client.post(
        "/api/v1/service-requests",
        json={"session_id": old_session_id, "request_type": "other", "message": "忘れ物確認"},
    ).json()
    message = client.post(
        "/api/v1/messages",
        json={"session_id": old_session_id, "sender_type": "kitchen", "text": "確認します"},
    ).json()

    status_response = client.get("/api/v1/admin/table-reset-status")
    assert status_response.headers["cache-control"] == "no-store"
    table_status = next(item for item in status_response.json() if item["table_id"] == 1)
    assert table_status["confirmed_total"] == confirmed["items"][0]["line_total"]
    assert table_status["draft_item_count"] == 2
    assert table_status["unserved_item_count"] == 0
    assert table_status["pending_request_count"] == 1
    assert table_status["can_reset"] is True

    with client.websocket_connect(f"/api/v1/ws/orders/{old_session_id}") as customer_ws:
        with client.websocket_connect("/api/v1/ws/kitchen") as kitchen_ws:
            reset = client.post(
                "/api/v1/admin/tables/1/reset",
                json={"expected_session_id": old_session_id},
            )
            assert reset.status_code == 200
            assert customer_ws.receive_json()["type"] == "table.session_reset"
            assert kitchen_ws.receive_json()["type"] == "table.session_reset"

    result = reset.json()
    assert result["closed_session_id"] == old_session_id
    assert result["new_session_id"] != old_session_id
    assert client.post(
        "/api/v1/admin/tables/1/reset",
        json={"expected_session_id": old_session_id},
    ).status_code == 409
    assert client.get(f"/api/v1/sessions/{old_session_id}/shared-order").status_code == 409
    assert client.get("/api/v1/kitchen/snapshot").json()["shared_orders"] == []

    with Session(engine) as db:
        old_session = db.get(TableSession, old_session_id)
        new_session = db.get(TableSession, result["new_session_id"])
        table = db.get(DiningTable, 1)
        order = db.get(Order, initial["id"])
        statuses = sorted((item.status, item.kitchen_status, item.quantity) for item in order.items)
        assert old_session.status == "closed"
        assert old_session.closed_at is not None
        assert new_session.status == "open"
        assert table.usage_status == "available"
        assert table.current_order_id is None
        assert order.status == "closed"
        assert ("cancelled", "cancelled", 2) in statuses
        assert ("confirmed", "served", 1) in statuses
        assert db.get(ServiceRequest, service_request["id"]).status == "cancelled"
        assert db.get(Message, message["id"]).dismissed_at is not None

    fresh_page = client.get("/table/1/order?lang=ja")
    assert fresh_page.status_code == 200
    assert f'data-session-id="{result["new_session_id"]}"' in fresh_page.text
    assert f'data-language-url="/table/1"' in fresh_page.text
    fresh_order = client.get(f"/api/v1/sessions/{result['new_session_id']}/shared-order").json()
    assert fresh_order["items"] == []
    assert fresh_order["current_total"] == 0


def test_admin_table_reset_is_blocked_while_confirmed_items_are_unserved(shared_client):
    client, engine = shared_client
    initial = client.get("/api/v1/tables/1/shared-order").json()
    menu_item_id = client.get("/api/v1/menu").json()[0]["id"]
    added = client.post(
        f"/api/v1/shared-orders/{initial['id']}/items",
        json={"menu_item_id": menu_item_id, "quantity": 2, "modifier_option_ids": [], "input_by": "staff", "input_language": "ja", "expected_version": initial["version"]},
    ).json()
    client.post(
        f"/api/v1/shared-orders/{initial['id']}/confirm",
        json={"expected_version": added["version"], "confirmed_by": "staff", "confirmation_method": "touch_button"},
    )

    table_status = next(item for item in client.get("/api/v1/admin/table-reset-status").json() if item["table_id"] == 1)
    assert table_status["unserved_item_count"] == 2
    assert table_status["can_reset"] is False
    assert table_status["block_reason"] == "未提供の商品があります"
    blocked = client.post(
        "/api/v1/admin/tables/1/reset",
        json={"expected_session_id": initial["session_id"]},
    )
    assert blocked.status_code == 409
    assert "未提供の商品が2点" in blocked.json()["detail"]
    with Session(engine) as db:
        assert db.get(TableSession, initial["session_id"]).status == "open"
        assert db.get(Order, initial["id"]).status == "open"


def test_admin_page_exposes_table_departure_reset_controls(shared_client):
    client, _ = shared_client
    page = client.get("/admin")
    assert page.status_code == 200
    assert page.headers["cache-control"] == "no-store"
    assert 'id="table-reset-list"' in page.text
    assert "テーブル管理" in page.text
    assert "退店処理して次のお客様へ" in page.text
    assert "table-admin.js?v=20260827-1" in page.text

    menu_page = client.get("/admin/menu")
    assert menu_page.status_code == 200
    assert "カテゴリ設定" in menu_page.text
    assert 'id="table-reset-list"' not in menu_page.text
    assert "admin.js?v=20260827-4" in menu_page.text


def test_kitchen_websocket_heartbeat_returns_pong(shared_client):
    client, _ = shared_client
    page = client.get("/kitchen")
    assert page.status_code == 200
    assert "kitchen.js?v=20260827-6" in page.text
    with client.websocket_connect("/api/v1/ws/kitchen") as websocket:
        websocket.send_text("ping")
        assert websocket.receive_json() == {"type": "connection.pong"}
