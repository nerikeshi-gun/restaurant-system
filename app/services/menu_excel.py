import csv
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.config import settings
from app.models import MenuCategory, MenuItem, MenuItemModifierGroup
from app.services.menu_catalog import import_menu_csv


HEADERS = [
    "商品コード", "商品名", "価格", "カテゴリ", "カテゴリ表示順", "販売中", "別名", "読み方",
    "短縮名", "表示順", "サイズ", "セット内容", "選択肢グループ1", "選択肢グループ2",
    "選択肢グループ3", "選択肢グループ4", "選択肢グループ5", "多言語表記",
]


def _groups(db: Session) -> list[MenuCategory]:
    return list(
        db.scalars(
            select(MenuCategory)
            .where(MenuCategory.is_option_group.is_(True), MenuCategory.is_active.is_(True))
            .options(selectinload(MenuCategory.options))
            .order_by(MenuCategory.display_order, MenuCategory.id)
        )
    )


def build_menu_workbook(db: Session, *, current: bool) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "メニュー"
    sheet.append(HEADERS)
    if current:
        items = list(
            db.scalars(
                select(MenuItem)
                .join(MenuCategory)
                .where(MenuItem.store_id == settings.default_store_id, MenuItem.is_archived.is_(False))
                .options(selectinload(MenuItem.category), selectinload(MenuItem.aliases), selectinload(MenuItem.modifier_groups).selectinload(MenuItemModifierGroup.group))
                .order_by(MenuCategory.display_order, MenuItem.display_order, MenuItem.id)
            )
        )
        for item in items:
            group_names = [link.group.name for link in sorted(item.modifier_groups, key=lambda value: (value.display_order, value.id)) if link.group.is_active][:5]
            sheet.append([
                item.menu_code or f"M{item.id:03d}", item.name, item.price, item.category.name,
                item.category.display_order, 1 if item.is_available else 0,
                "|".join(alias.alias for alias in item.aliases), item.reading, item.short_name,
                item.display_order, item.size_label, item.set_contents, *group_names,
                *([""] * (5 - len(group_names))), item.localized_names_json or "{}",
            ])
    else:
        sheet.append(["M001", "チキンカレー", 1100, "カレー", 10, 1, "チキン|チキンカレー", "ちきんかれー", "チキン", 10, "", "サラダ付き", "", "", "", "", "", '{"en":"Chicken Curry"}'])
        sheet.append(["M002", "マトンマサラ", 1000, "カレー", 10, 1, "マトン", "まとんまさら", "マトン", 20, "", "", "", "", "", "", "", '{"en":"Mutton Masala"}'])

    groups = _groups(db)
    master = workbook.create_sheet("選択肢マスタ")
    master.append(["グループ名", "選び方", "登録済みの選択肢"])
    for group in groups:
        names = [option.name for option in sorted(group.options, key=lambda value: (value.display_order, value.id)) if option.is_available]
        master.append([group.name, "1つ選択" if group.selection_type == "single" else "複数選択", "、".join(names)])
    if not groups:
        master.append(["", "", ""])

    workbook.defined_names.add(DefinedName("ToppingGroups", attr_text=f"'選択肢マスタ'!$A$2:$A${max(2, len(groups) + 1)}"))
    validation = DataValidation(type="list", formula1="=ToppingGroups", allow_blank=True)
    validation.error = "選択肢マスタに登録されているグループを選んでください"
    validation.errorTitle = "選択できないグループです"
    validation.prompt = "▼から選択肢グループを選びます"
    validation.promptTitle = "選択肢グループ"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add("M2:Q1001")

    header_fill = PatternFill("solid", fgColor="7A2E1D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(2, sheet.max_row)}"
    widths = [14, 24, 11, 16, 14, 10, 24, 20, 18, 10, 12, 24, 19, 19, 19, 19, 19, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=max(2, sheet.max_row), min_col=3, max_col=3):
        row[0].number_format = '#,##0'

    for cell in master[1]:
        cell.fill = PatternFill("solid", fgColor="315F86")
        cell.font = Font(color="FFFFFF", bold=True)
    master.column_dimensions["A"].width = 24
    master.column_dimensions["B"].width = 14
    master.column_dimensions["C"].width = 50
    master.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def import_menu_workbook(db: Session, raw: bytes) -> dict:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if "メニュー" not in workbook.sheetnames:
        raise ValueError("『メニュー』シートがありません")
    sheet = workbook["メニュー"]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError("メニューがありません")
    headers = [str(value or "").strip() for value in rows[0]]
    index = {name: position for position, name in enumerate(headers)}
    required = {"商品コード", "商品名", "価格", "カテゴリ"}
    if not required.issubset(index):
        raise ValueError("商品コード、商品名、価格、カテゴリの列が必要です")
    groups = {group.name: group for group in _groups(db)}
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["menu_id", "name_ja", "price", "category", "category_order", "available", "alias", "reading", "short_name", "display_order", "size", "set_contents", "toppings", "localized_names"])

    def value(row: tuple, name: str, default=""):
        position = index.get(name)
        return default if position is None or position >= len(row) or row[position] is None else row[position]

    item_count = 0
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(cell not in (None, "") for cell in row):
            continue
        selected_groups = []
        for number in range(1, 6):
            group_name = str(value(row, f"選択肢グループ{number}")).strip()
            if not group_name:
                continue
            group = groups.get(group_name)
            if not group:
                raise ValueError(f"{row_number}行目の選択肢グループ『{group_name}』が見つかりません")
            options = [option.name for option in sorted(group.options, key=lambda item: (item.display_order, item.id)) if option.is_available]
            if options:
                selected_groups.append(",".join([group.name, *options]))
        writer.writerow([
            value(row, "商品コード"), value(row, "商品名"), value(row, "価格"), value(row, "カテゴリ"),
            value(row, "カテゴリ表示順", 10), value(row, "販売中", 1), value(row, "別名"), value(row, "読み方"),
            value(row, "短縮名"), value(row, "表示順", row_number * 10), value(row, "サイズ"), value(row, "セット内容"),
            " / ".join(selected_groups), value(row, "多言語表記", "{}"),
        ])
        item_count += 1
    if not item_count:
        raise ValueError("メニューがありません")
    return import_menu_csv(db, ("\ufeff" + output.getvalue()).encode("utf-8"))
